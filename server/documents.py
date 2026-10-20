"""Local Office inspection, rendering, and verified copy editing. Never executes macros."""
import sys, json, zipfile, os, re, copy, subprocess, hashlib, shutil
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from openpyxl import load_workbook
from lxml import etree

NS={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
def validate_zip(path):
 with zipfile.ZipFile(path) as z:
  if len(z.infolist())>20000 or sum(i.file_size for i in z.infolist())>200_000_000: raise ValueError('Officeファイルの展開サイズが上限を超えています')
  if any('..' in Path(i.filename).parts or i.filename.startswith('/') for i in z.infolist()): raise ValueError('不正なOfficeファイルです')
  if any('vbaProject' in i.filename for i in z.infolist()): raise ValueError('マクロ付きファイルには対応していません')
def all_shapes(shapes):
 for s in shapes:
  yield s
  if s.shape_type==MSO_SHAPE_TYPE.GROUP: yield from all_shapes(s.shapes)
def color_info(color):
 try: return str(color.rgb) if color.type else None
 except (AttributeError,ValueError): return None
def extract(path):
 validate_zip(path)
 if path.suffix.lower()=='.xlsx':
  wb=load_workbook(path,read_only=True,data_only=True)
  sheets=[]
  total=0
  for ws in wb:
   rows=[]
   for i,row in enumerate(ws.iter_rows(values_only=True),1):
    if any(v is not None for v in row):
     total+=1
     if total>10000: raise ValueError('Excelは合計1万行以下にしてください')
     cells=list(row)
     while cells and cells[-1] is None:cells.pop()
     if len(cells)>100:raise ValueError('Excelは100列以下にしてください。超過部分を省略せず停止しました')
     rows.append({'row':i,'cells':[str(v) if v is not None else '' for v in cells]})
   sheets.append({'name':ws.title,'rows':rows})
  return {'type':'xlsx','sheets':sheets}
 p=Presentation(path)
 if len(p.slides)>80: raise ValueError('PPTXは80スライド以下にしてください')
 slides=[]
 for n,slide in enumerate(p.slides,1):
  shapes=[]
  for shape in all_shapes(slide.shapes):
   text=shape.text if shape.has_text_frame else ''
   runs=[]
   if shape.has_text_frame:
    for para in shape.text_frame.paragraphs:
     for run in para.runs:
      runs.append({'text':run.text,'color':color_info(run.font.color)})
   tables=[]
   if shape.has_table:
    tables=[['\n'.join(par.text for par in cell.text_frame.paragraphs) for cell in row.cells] for row in shape.table.rows]
   try: fill=color_info(shape.fill.fore_color) if shape.fill.type else None
   except (AttributeError,ValueError,TypeError):fill=None
   try: line=color_info(shape.line.color)
   except (AttributeError,ValueError,TypeError):line=None
   shapes.append({'id':shape.shape_id,'name':shape.name,'text':text,'x':shape.left,'y':shape.top,'width':shape.width,'height':shape.height,'runs':runs,'fill':fill,'line':line,'shapeType':str(shape.shape_type),'table':tables})
  notes=slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ''
  slides.append({'number':n,'shapes':shapes,'notes':notes})
 comments=[]
 with zipfile.ZipFile(path) as z:
  for name in z.namelist():
   if 'comment' in name.lower() and name.endswith('.xml'):
    root=etree.fromstring(z.read(name))
    comments.append({'part':name,'text':' '.join(root.itertext())})
 return {'type':'pptx','slideWidth':p.slide_width,'slideHeight':p.slide_height,'slides':slides,'comments':comments}

def render(path,directory):
 directory.mkdir(exist_ok=True)
 pdf=directory/(path.stem+'.pdf')
 if not pdf.exists():
  profile=directory/'lo-profile'
  result=subprocess.run(['soffice','--headless','-env:UserInstallation='+profile.as_uri(),'--convert-to','pdf','--outdir',str(directory),str(path)],capture_output=True,text=True,timeout=150)
  if not pdf.exists():raise ValueError('スライド画像化に失敗しました。LibreOfficeの導入状態とファイルを確認してください')
 import pymupdf as fitz
 doc=fitz.open(pdf)
 for i,page in enumerate(doc,1):
  img=directory/f'slide-{i}.jpg'
  if not img.exists():page.get_pixmap(matrix=fitz.Matrix(1.1,1.1),alpha=False).save(img,jpg_quality=78)
 return len(doc)

def replace_in_frame(frame,old,new):
 # Cross-run replacement retains the formatting of the first affected run.
 for para in frame.paragraphs:
  text=''.join(run.text for run in para.runs)
  if old not in text:continue
  if text.count(old)!=1:raise ValueError('置換対象が複数あり曖昧です。対象を絞ってください')
  start=text.index(old);end=start+len(old);offset=0;inserted=False
  for run in para.runs:
   value=run.text;a=offset;b=offset+len(value);offset=b
   if b<=start or a>=end:continue
   left=value[:max(0,start-a)] if a<start else ''
   right=value[max(0,end-a):] if b>end else ''
   run.text=left+(new if not inserted else '')+right;inserted=True
  return
 # Multi-paragraph replacements must retain paragraph count to preserve structure.
 if frame.text==old and len(old.split('\n'))==len(new.split('\n')):
  for para,value in zip(frame.paragraphs,new.split('\n')):
   if para.runs:
    para.runs[0].text=value
    for r in para.runs[1:]:r.text=''
   else:para.text=value
  return
 raise ValueError('対象本文と一致しないか、段落をまたぐ変更です。修正対象を段落単位で指定してください')

def strip_comments(path):
 # Remove comment parts, relationships, and content type declarations consistently.
 tmp=path.with_suffix('.clean.pptx')
 with zipfile.ZipFile(path) as src,zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as dst:
  removed={n for n in src.namelist() if 'comment' in n.lower() and (n.startswith('ppt/') or n.startswith('docProps/'))}
  for name in src.namelist():
   if name in removed:continue
   raw=src.read(name)
   if name.endswith('.rels') or name=='[Content_Types].xml':
    root=etree.fromstring(raw)
    for el in list(root):
     if 'comment' in el.get('Type','').lower() or 'comment' in el.get('PartName','').lower() or 'comment' in el.get('Target','').lower():root.remove(el)
    raw=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
   elif re.match(r'ppt/slides/slide\d+\.xml$',name):
    root=etree.fromstring(raw)
    for el in root.xpath('//*[contains(local-name(), "cmLst") or contains(local-name(), "comment")]'):
     if el.getparent() is not None:el.getparent().remove(el)
    raw=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
   dst.writestr(name,raw)
 os.replace(tmp,path)

def apply_revision(source,output,plan):
 p=Presentation(source); original_count=len(p.slides);applied=[];deleted=[]
 before={(n,s.shape_id):{'text':s.text if s.has_text_frame else None,'geometry':(s.left,s.top,s.width,s.height)} for n,slide in enumerate(p.slides,1) for s in all_shapes(slide.shapes)}
 for item in plan['annotations']:
  if not item.get('approved'):continue
  if item['status']!='applicable':raise ValueError('要確認の指摘は承認できません')
  n=item['slide'];slide=p.slides[n-1];shapes={s.shape_id:s for s in all_shapes(slide.shapes)}
  for change in item['replacements']:
   shape=shapes.get(change['shapeId'])
   if shape is None or not shape.has_text_frame:raise ValueError('本文のテキストボックスが見つかりません')
   replace_in_frame(shape.text_frame,change['oldText'],change['newText']);applied.append({'slide':n,**change})
 for item in plan['annotations']:
  if not item.get('approved'):continue
  slide=p.slides[item['slide']-1];protected={r['shapeId'] for x in plan['annotations'] for r in x['replacements'] if x['slide']==item['slide']}
  for shape in list(all_shapes(slide.shapes)):
   if shape.shape_id in item['annotationShapeIds']:
    if shape.shape_id in protected:raise ValueError('本文と注釈の削除指定が重複しています')
    el=shape._element;el.getparent().remove(el);deleted.append({'slide':item['slide'],'shapeId':shape.shape_id})
 # Existing notes are kept verbatim. Only recognized review shapes and Office comments are removed.
 p.save(output)
 if plan.get('removeOfficeComments'):strip_comments(output)
 check=Presentation(output)
 if len(check.slides)!=original_count:raise ValueError('スライド枚数の検証に失敗しました')
 for item in applied:
  shapes={s.shape_id:s for s in all_shapes(check.slides[item['slide']-1].shapes)}
  if item['newText'] not in shapes[item['shapeId']].text:raise ValueError('修正本文の再読込検証に失敗しました')
 for item in deleted:
  if any(s.shape_id==item['shapeId'] for s in all_shapes(check.slides[item['slide']-1].shapes)):raise ValueError('注釈削除の検証に失敗しました')
 touched={(r['slide'],r['shapeId']) for r in applied};removed={(r['slide'],r['shapeId']) for r in deleted}
 after={(n,s.shape_id):s for n,slide in enumerate(check.slides,1) for s in all_shapes(slide.shapes)}
 for identity,original in before.items():
  if identity in removed:continue
  shape=after.get(identity)
  if shape is None:raise ValueError('指定以外の図形が削除されたため検証に失敗しました')
  if (shape.left,shape.top,shape.width,shape.height)!=original['geometry']:raise ValueError('元の図形位置・サイズの保持検証に失敗しました')
  if identity not in touched and original['text'] is not None and shape.text!=original['text']:raise ValueError('指定以外の本文が変わったため検証に失敗しました')
 validate_zip(output)
 return {'applied':len(applied),'deletedShapes':len(deleted),'slides':len(check.slides),'sha256':hashlib.sha256(output.read_bytes()).hexdigest(),'verification':'本文置換・選択注釈削除・未変更本文と図形位置の保持・スライド枚数・再読込を検証。目視による最終確認が必要。'}

def clone_x(source,output,plan):
 p=Presentation(source); idx=plan['templateSlide']-1
 if idx<0 or idx>=len(p.slides):raise ValueError('テンプレートのスライド番号が不正です')
 template=p.slides[idx]; original=len(p.slides)
 for item in plan['ideas']:
  slide=p.slides.add_slide(template.slide_layout)
  for s in list(slide.shapes):slide.shapes._spTree.remove(s._element)
  rid_map={}
  for rel in template.part.rels.values():
   if rel.reltype.endswith('/notesSlide') or rel.reltype.endswith('/slideLayout') or 'comment' in rel.reltype.lower():continue
   rid_map[rel.rId]=slide.part.rels.get_or_add_ext_rel(rel.reltype,rel.target_ref) if rel.is_external else slide.part.rels.get_or_add(rel.reltype,rel.target_part)
  for shape in template.shapes:
   el=copy.deepcopy(shape.element)
   for child in el.iter():
    for name,value in list(child.attrib.items()):
     if name.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}') and value in rid_map:child.set(name,rid_map[value])
   slide.shapes._spTree.insert_element_before(el,'p:extLst')
  if template._element.cSld.bg is not None:slide._element.cSld.insert(0,copy.deepcopy(template._element.cSld.bg))
  shapes={s.shape_id:s for s in all_shapes(slide.shapes)}
  for field,shape_id in plan['mapping'].items():
   if not shape_id:continue
   shape=shapes.get(int(shape_id))
   if shape is None or not shape.has_text_frame:raise ValueError('指定したテキストボックスが見つかりません')
   value=item.get(field,'')
   # Keep first-run font and shape geometry; additional paragraphs are not inherited.
   frame=shape.text_frame;para=frame.paragraphs[0]
   for other in list(frame.paragraphs)[1:]:other._p.getparent().remove(other._p)
   if para.runs:
    para.runs[0].text=value
    for run in para.runs[1:]:run.text=''
   else:para.text=value
  slide.notes_slide.notes_text_frame.text='AI生成ドラフト / 承認前\n出典: '+ '\n'.join(item.get('sourceUrls',[]))+'\n確認事項: '+item.get('caution','')+'\n画像ラフ: '+item.get('visual','')+'\nクイズの答え: '+item.get('answer','')
 # Keep the input deck intact in the output, append editable review slides.
 p.save(output);check=Presentation(output)
 if len(check.slides)!=original+len(plan['ideas']):raise ValueError('企画スライド数の検証に失敗しました')
 for slide,idea in zip(list(check.slides)[original:],plan['ideas']):
  shape_map={s.shape_id:s for s in all_shapes(slide.shapes)}
  for field,sid in plan['mapping'].items():
   if sid and idea.get(field,'') not in shape_map[int(sid)].text:raise ValueError('企画本文の再読込検証に失敗しました')
 return {'slides':len(check.slides),'appendedSlides':len(plan['ideas']),'sha256':hashlib.sha256(output.read_bytes()).hexdigest(),'verification':'元資料を保持し、指定スライドを複製・本文置換して追加。再読込検証済み。レイアウトの目視確認が必要。'}

if __name__=='__main__':
 try:
  cmd=sys.argv[1];path=Path(sys.argv[2]).resolve()
  if cmd=='extract':result=extract(path)
  elif cmd=='render':result={'slides':render(path,Path(sys.argv[3]).resolve())}
  elif cmd=='revise':result=apply_revision(path,Path(sys.argv[3]),json.load(sys.stdin))
  elif cmd=='clone-x':result=clone_x(path,Path(sys.argv[3]),json.load(sys.stdin))
  else:raise ValueError('未知の操作です')
  print(json.dumps(result,ensure_ascii=False,default=str))
 except Exception as e:
  print(json.dumps({'error':str(e)},ensure_ascii=False));sys.exit(1)
