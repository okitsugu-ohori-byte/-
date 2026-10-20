from pathlib import Path
import sys
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from openpyxl import Workbook

root=Path(sys.argv[1]);root.mkdir(parents=True,exist_ok=True)
p=Presentation();p.slide_width=Inches(13.333);p.slide_height=Inches(7.5)
s=p.slides.add_slide(p.slide_layouts[6])
fields=[('title','秋の家事を、もっと快適に。',.6,.5,12,.7),('body','本日は洗濯の工夫をご紹介します。',.6,1.6,7,1.2),('reply','詳細はこちら。',.6,3.1,7,.8),('claim','商品情報は承認資料を確認。',.6,4.3,7,.8),('visual','画像ラフ：洗濯物と暮らしのシーン',8,1.6,4.5,3.5)]
for name,text,x,y,w,h in fields:
 shape=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h));shape.name=name
 para=shape.text_frame.paragraphs[0]
 if name=='body':
  para.add_run().text='本日は洗濯の';para.add_run().text='工夫をご紹介します。'
 else:para.add_run().text=text
 for run in para.runs:
  run.font.name='Noto Sans CJK JP';run.font.size=Pt(18);run.font.color.rgb=RGBColor.from_string('344638')
note=s.shapes.add_textbox(Inches(8), Inches(5.5), Inches(4.6), Inches(.8));note.name='review-note';note.text='「本日は」を「今日は」に修正';note.text_frame.paragraphs[0].runs[0].font.color.rgb=RGBColor(255,0,0)
s.notes_slide.notes_text_frame.text='審査用の元ノート'
p.save(root/'review.pptx')
wb=Workbook();ws=wb.active;ws.title='10月';ws.append(['SNS','テーマ','商品','注意点']);ws.append(['X','季節の家事','洗濯機','商品性能は要確認']);wb.save(root/'SNS企画案_10月.xlsx')
print(str(root/'review.pptx'))
